"""Deterministic GOST DOCX; PDF is converted from that exact DOCX by LibreOffice."""

import json
import subprocess
import tempfile
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader

from .config import Settings
from .models import Report
from .validator import ValidationFailed, normalized, validate_docx, validate_pdf


def bibliography(source):
    author = source.authors[0]
    authors = ", ".join(source.authors[:3]) + (" [и др.]" if len(source.authors) > 3 else "")
    venue = source.venue or ("arXiv" if source.arxiv_id else "Электронная публикация")
    doi = f" - DOI: {source.doi}." if source.doi else ""
    # Missing volume/pages are never invented. Electronic-resource description is used.
    return (
        f"{author}. {source.title} [Электронный ресурс] / {authors} // {venue}. - {source.year}."
        f"{doi} - URL: {source.url} (дата обращения: {date.today():%d.%m.%Y})."
    )


def number_sources(report: Report) -> Report:
    order = []
    for section in report.sections:
        for paragraph in section.paragraphs:
            for sid in paragraph.source_ids:
                if sid not in order:
                    order.append(sid)
    mapping = {old: new for new, old in enumerate(order, 1)}
    result = report.model_copy(deep=True)
    known = {s.id: s for s in result.sources}
    result.sources = [known[old] for old in order]
    for old, new in mapping.items():
        known[old].id = new
    for section in result.sections:
        for paragraph in section.paragraphs:
            paragraph.source_ids = [mapping[sid] for sid in paragraph.source_ids]
            paragraph.evidence_quotes = {
                str(mapping[int(k)]): v for k, v in paragraph.evidence_quotes.items() if int(k) in mapping
            }
    return result


def heading_labels(report):
    labels = []
    number = 0
    for section in report.sections:
        if section.title in ("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ"):
            labels.append(section.title)
        else:
            number += 1
            labels.append(f"{number} {section.title}")
    return labels + ["СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ПРИЛОЖЕНИЕ А"]


class Reporter:
    def __init__(self, settings: Settings):
        self.settings = settings

    def chart(self, report: Report, directory: Path):
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        counts = Counter(s.year for s in report.sources)
        fig, ax = plt.subplots(figsize=(6.2, 2.2), layout="constrained")
        ax.bar([str(y) for y in sorted(counts)], [counts[y] for y in sorted(counts)], color="#555555")
        ax.set_xlabel("Год публикации")
        ax.set_ylabel("Число источников")
        ax.yaxis.get_major_locator().set_params(integer=True)
        ax.spines[["top", "right"]].set_visible(False)
        path = directory / "sources_by_year.png"
        fig.savefig(path, dpi=180)
        plt.close(fig)
        return path

    def document(self, report: Report, path: Path, chart: Path, pages=None, total_pages=0):
        pages = pages or {}
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
        sec.footer_distance = Mm(10)
        sec.different_first_page_header_footer = True
        for name in ("Normal", "Title", "Heading 1", "Caption"):
            style = doc.styles[name]
            style.font.name, style.font.size, style.font.color.rgb = (
                "Times New Roman",
                Pt(14),
                RGBColor(0, 0, 0),
            )
            fonts = style.element.get_or_add_rPr().rFonts
            for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + attribute), None)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + attribute), "Times New Roman")
            rpr = style.element.get_or_add_rPr()
            size_cs = rpr.find(qn("w:szCs"))
            if size_cs is None:
                size_cs = OxmlElement("w:szCs")
                rpr.append(size_cs)
            size_cs.set(qn("w:val"), "28")
            for border in style.element.xpath(".//w:pBdr"):
                border.getparent().remove(border)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
        normal = doc.styles["Normal"]
        normal.font.bold = False
        doc.styles["Caption"].font.bold = False
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.first_line_indent = Mm(12.5)
        normal.paragraph_format.widow_control = True
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.styles["Title"].font.bold = True
        doc.styles["Heading 1"].font.bold = True
        doc.styles["Heading 1"].paragraph_format.keep_with_next = True
        doc.styles["Heading 1"].paragraph_format.space_after = Pt(14)
        doc.styles["Caption"].paragraph_format.line_spacing = 1
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Mm(0)
        # Center on the sheet, accounting for asymmetric 30/15 mm text margins.
        footer.paragraph_format.left_indent = Mm(-15)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        cached = OxmlElement("w:t")
        cached.text = "1"
        run.append(cached)
        field.append(run)
        footer._p.append(field)

        def centered(text, style=None):
            p = doc.add_paragraph(text, style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Mm(0)
            return p

        brief = report.brief
        centered(brief["organization"])
        centered(brief.get("department", "Учебная исследовательская работа"))
        doc.add_paragraph()
        centered("ОТЧЕТ О НАУЧНО-ИССЛЕДОВАТЕЛЬСКОЙ РАБОТЕ", "Title")
        centered(report.title, "Title")
        centered("Заключительный отчет")
        doc.add_paragraph()
        doc.add_paragraph("Исполнитель: " + brief["author"])
        doc.add_paragraph("Руководитель: " + brief["supervisor"])
        if brief.get("registration"):
            doc.add_paragraph(brief["registration"])
        doc.add_paragraph()
        centered(f"{brief['city']} {date.today().year}")

        def heading(label, structural=True):
            p = doc.add_paragraph(label, "Heading 1")
            p.paragraph_format.page_break_before = True
            p.paragraph_format.first_line_indent = Mm(0) if structural else Mm(12.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if structural else WD_ALIGN_PARAGRAPH.LEFT
            return p

        heading("РЕФЕРАТ")
        doc.add_paragraph(
            f"Отчет {total_pages or '-'} с., 1 рис., {1 if report.metrics else 0} табл., "
            f"{len(report.sources)} источн., 1 прил."
        )
        doc.add_paragraph(", ".join(report.keywords).upper())
        doc.add_paragraph(report.abstract)
        heading("СОДЕРЖАНИЕ")
        labels = heading_labels(report)
        for label in labels:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(165), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            p.add_run(label + "\t" + str(pages.get(label, "-")))
        for section, label in zip(report.sections, labels):
            heading(label, section.title in ("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ"))
            for paragraph in section.paragraphs:
                suffix = (
                    " [" + ", ".join(str(sid) for sid in paragraph.source_ids) + "]"
                    if paragraph.source_ids
                    else ""
                )
                doc.add_paragraph(paragraph.text.rstrip() + suffix)
        heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
        for source in report.sources:
            p = doc.add_paragraph(f"{source.id} {bibliography(source)}")
            p.paragraph_format.first_line_indent = Mm(0)
        heading("ПРИЛОЖЕНИЕ А")
        centered("Характеристики корпуса источников и данных")
        doc.add_paragraph(
            "В анализ включены только найденные во внешних научных базах записи. "
            "Источники с доступной аннотацией исследованы в пределах этой аннотации; "
            "наличие записи в базе и цитируемость не доказывают истинность выводов. "
            "Распределение включенных публикаций приведено на рисунке А.1."
        )
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.add_run().add_picture(str(chart), width=Mm(155))
        p.paragraph_format.keep_with_next = True
        centered("Рисунок А.1 - Распределение включенных источников по годам", "Caption")
        for source in report.sources:
            level = (
                "полный текст (извлечен частично)" if source.evidence_level == "full_text" else "аннотация"
            )
            doc.add_paragraph(
                f"Источник {source.id}: {source.provider}, {level}. " + "; ".join(source.quality_notes)
            )
        if report.metrics:
            doc.add_paragraph(
                f"Проанализировано строк CSV: {report.metrics['rows']}. "
                f"Контрольная сумма SHA-256: {report.metrics['sha256']}. "
                "Описательные статистики представлены в таблице А.1. "
                "Пропуски и нечисловые значения исключались отдельно для каждого столбца."
            )
            p = doc.add_paragraph("Таблица А.1 - Описательные статистики числовых столбцов")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Mm(0)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ("Столбец", "n", "Среднее", "Медиана", "Станд откл")):
                cell.text = text
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for name, stats in report.metrics["columns"].items():
                cells = table.add_row().cells
                for cell, value in zip(
                    cells, (name, stats["n"], stats["mean"], stats["median"], stats["std"])
                ):
                    cell.text = f"{value:.4g}" if isinstance(value, float) else str(value)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.first_line_indent = Mm(0)
                        p.paragraph_format.line_spacing = 1
                        for run in p.runs:
                            run.font.size = Pt(12)
        doc.core_properties.title = report.title
        doc.core_properties.author = brief["author"]
        doc.core_properties.subject = "Научный обзор и анализ предоставленных данных"
        doc.save(path)

    def convert(self, docx: Path, directory: Path):
        if not self.settings.soffice:
            raise RuntimeError(
                "Установите LibreOffice и задайте SOFFICE_PATH: строгий PDF создается из DOCX."
            )
        pdf = directory / (docx.stem + ".pdf")
        if pdf.exists():
            pdf.unlink()
        with tempfile.TemporaryDirectory(prefix="gost-lo-") as profile:
            command = [
                self.settings.soffice,
                "-env:UserInstallation=" + Path(profile).as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(directory),
                str(docx),
            ]
            result = subprocess.run(command, capture_output=True, timeout=120, check=False)
            if result.returncode or not pdf.exists():
                raise RuntimeError("LibreOffice не создал PDF. Проверьте SOFFICE_PATH и установку шрифтов.")
        return pdf

    def export(self, report: Report, directory: Path) -> dict:
        directory.mkdir(parents=True, exist_ok=True)
        report = number_sources(report)
        chart = self.chart(report, directory)
        path = directory / "research_report.docx"
        pages, total = {}, 0
        for _ in range(4):
            self.document(report, path, chart, pages, total)
            pdf = self.convert(path, directory)
            reader = PdfReader(pdf)
            found = {}
            for i, page in enumerate(reader.pages, 1):
                lines = (page.extract_text() or "").splitlines()
                for label in heading_labels(report):
                    if any(normalized(line) == normalized(label) for line in lines):
                        found[label] = i  # TOC has page numbers; exact heading matches body only.
            if found == pages and len(reader.pages) == total:
                break
            pages, total = found, len(reader.pages)
        else:
            raise ValidationFailed("Пагинация содержания не стабилизировалась")
        if len(pages) != len(heading_labels(report)):
            raise ValidationFailed("Не удалось установить страницы всех разделов")
        docx_checks = validate_docx(path)
        pdf_checks = validate_pdf(pdf, ["РЕФЕРАТ", "СОДЕРЖАНИЕ"] + heading_labels(report))
        quality = {
            "status": "passed_automated_checks",
            "standard": "ГОСТ 7.32-2017",
            "bibliography_standard": "ГОСТ 7.1-2003, электронные ресурсы",
            "docx": docx_checks,
            "pdf": pdf_checks,
            "contents": pages,
            "sources": len(report.sources),
            "limitations": [
                "Автоматическая проверка не заменяет визуальный нормоконтроль",
                "Университетский титульный лист; регистрационные сведения только из ТЗ",
                "Self-check одной LLM не гарантирует истинность каждого утверждения",
            ],
        }
        (directory / "quality.json").write_text(
            json.dumps(quality, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (directory / "report.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
        (directory / "sources.json").write_text(
            json.dumps([s.model_dump() for s in report.sources], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (directory / "evidence.json").write_text(
            json.dumps(
                [
                    {"section": section.title, **paragraph.model_dump()}
                    for section in report.sections
                    for paragraph in section.paragraphs
                ],
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return {
            "docx": str(path),
            "pdf": str(pdf),
            "quality": str(directory / "quality.json"),
            "report": str(directory / "report.json"),
            "sources": str(directory / "sources.json"),
            "evidence": str(directory / "evidence.json"),
        }
