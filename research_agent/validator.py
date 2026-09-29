import re
from pathlib import Path

from docx import Document
from docx.shared import Mm
from pypdf import PdfReader

from .models import Report, Section, Source

SECTION_TITLES = [
    "ВВЕДЕНИЕ",
    "Обзор литературы",
    "Методология исследования",
    "Результаты и обсуждение",
    "Выводы и рекомендации",
    "ЗАКЛЮЧЕНИЕ",
]


class ValidationFailed(RuntimeError):
    pass


def normalized(text):
    return re.sub(r"\s+", " ", text).strip().casefold()


def validate_section(section: Section, sources: list[Source]) -> list[str]:
    known = {s.id: s for s in sources}
    errors = []
    for index, paragraph in enumerate(section.paragraphs, 1):
        # Confidence intervals such as [0,546; 0,799] are data, not bibliography citations.
        if re.search(r"\[[1-9]\d*(?:, [1-9]\d*)*\]", paragraph.text):
            errors.append(f"Абзац {index}: модель вставила ссылку в текст вместо source_ids")
        if paragraph.kind == "evidence" and not paragraph.source_ids:
            errors.append(f"Абзац {index}: факт без источника")
        for sid in paragraph.source_ids:
            if sid not in known:
                errors.append(f"Абзац {index}: неизвестный источник {sid}")
                continue
            quote = paragraph.evidence_quotes.get(str(sid), "")
            if paragraph.kind == "evidence" and not 20 <= len(quote) <= 220:
                errors.append(f"Абзац {index}: нет достаточного фрагмента источника {sid}")
            if quote and normalized(quote) not in normalized(known[sid].text):
                errors.append(f"Абзац {index}: цитата отсутствует в источнике {sid}")
        if len(set(paragraph.source_ids)) != len(paragraph.source_ids):
            errors.append(f"Абзац {index}: повтор id")
    return errors


def validate_report(report: Report, minimum=15):
    errors = []
    if [s.title for s in report.sections] != SECTION_TITLES:
        errors.append("Нарушена структура разделов")
    if len(report.sources) < minimum:
        errors.append("Меньше 15 источников")
    if any(not s.metadata_verified for s in report.sources):
        errors.append("Метаданные источника не проверены внешним API")
    used = set()
    for section in report.sections:
        errors.extend(validate_section(section, report.sources))
        used.update(sid for p in section.paragraphs for sid in p.source_ids)
    if len(used) < minimum:
        errors.append(f"В тексте процитировано {len(used)} источников; требуется минимум {minimum}")
    if errors:
        raise ValidationFailed("; ".join(errors))


def validate_docx(path: Path):
    doc = Document(path)
    errors = []
    for section in doc.sections:
        for name, mm in (
            ("page_width", 210),
            ("page_height", 297),
            ("left_margin", 30),
            ("right_margin", 15),
            ("top_margin", 20),
            ("bottom_margin", 20),
        ):
            if abs(getattr(section, name) - Mm(mm)) > Mm(0.1):
                errors.append(f"DOCX: {name}")
        if not section.different_first_page_header_footer:
            errors.append("DOCX: титульный лист не исключен из видимой нумерации")
    style = doc.styles["Normal"]
    for name in ("Title", "Heading 1", "Caption"):
        current = doc.styles[name]
        if current.font.name != "Times New Roman" or current.element.xpath(".//w:pBdr"):
            errors.append(f"DOCX: неверный стиль {name}")
        fonts = current.element.get_or_add_rPr().rFonts
        if fonts is not None and any("Theme" in key or "theme" in key for key in fonts.attrib):
            errors.append(f"DOCX: шрифт темы переопределяет {name}")
    if style.font.name != "Times New Roman" or style.font.size.pt != 14:
        errors.append("DOCX: шрифт основного текста")
    if style.paragraph_format.line_spacing != 1.5:
        errors.append("DOCX: интервал")
    if abs(style.paragraph_format.first_line_indent - Mm(12.5)) > Mm(0.1):
        errors.append("DOCX: отступ")
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal" and paragraph.text:
            for run in paragraph.runs:
                if run.font.size and run.font.size.pt < 12:
                    errors.append("DOCX: основной текст меньше 12 пт")
    if errors:
        raise ValidationFailed("; ".join(errors))
    return {
        "a4": True,
        "margins_mm": [30, 15, 20, 20],
        "font": "Times New Roman 14",
        "line_spacing": 1.5,
        "first_line_mm": 12.5,
    }


def validate_pdf(path: Path, expected_headings: list[str]):
    reader = PdfReader(path)
    text = "\n".join(p.extract_text() or "" for p in reader.pages)
    errors = []
    for page in reader.pages:
        if abs(float(page.mediabox.width) - 595.28) > 2 or abs(float(page.mediabox.height) - 841.89) > 2:
            errors.append("PDF: формат не А4")
    for title in expected_headings:
        if normalized(title) not in normalized(text):
            errors.append(f"PDF: нет раздела {title}")
    if len(text) < 1000 or "\ufffd" in text:
        errors.append("PDF: текст не извлекается корректно")
    # Verify visible footer page numbers; title is page 1 without a number.
    for i, page in enumerate(reader.pages):
        footer_text = []

        def visitor(value, cm, tm, font, font_size):
            y = cm[1] * tm[4] + cm[3] * tm[5] + cm[5]
            if 0 <= y <= 60:
                footer_text.extend(value.strip().splitlines())

        page.extract_text(visitor_text=visitor)
        if i and str(i + 1) not in [line.strip() for line in footer_text]:
            errors.append(f"PDF: нет сквозного номера {i + 1} внизу страницы")
    if errors:
        raise ValidationFailed("; ".join(errors))
    return {"pages": len(reader.pages), "extractable_text": True, "a4": True, "page_numbers": True}
