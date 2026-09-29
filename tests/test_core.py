import httpx
import pytest

from research_agent.analysis import analyze_csv
from research_agent.config import Settings
from research_agent.reporter import bibliography, number_sources
from research_agent.retriever import Corpus, Retriever, SearchUnavailable, deduplicate
from research_agent.storage import Store
from research_agent.validator import ValidationFailed, validate_docx, validate_report, validate_section


def test_dedup_title_doi_arxiv(sources):
    same = sources[0].model_copy(update={"title": "different", "doi": "10.test/same"})
    other = sources[1].model_copy(update={"doi": "10.TEST/SAME"})
    assert len(deduplicate([sources[0], sources[0], same, other])) == 2


def test_fake_quote_rejected(report):
    report.sections[0].paragraphs[0].evidence_quotes = {"1": "invented experiment result"}
    assert "цитата отсутствует" in " ".join(validate_section(report.sections[0], report.sources))


def test_confidence_interval_not_mistaken_for_citation(report):
    report.sections[0].paragraphs[0].text += " 95% ДИ [0,546; 0,799]."
    assert validate_section(report.sections[0], report.sources) == []


def test_unknown_citation_rejected(report):
    report.sections[0].paragraphs[0].source_ids = [999]
    with pytest.raises(ValidationFailed, match="неизвестный источник"):
        validate_report(report)


def test_less_than_15_cited_rejected(report):
    for section in report.sections:
        for paragraph in section.paragraphs:
            paragraph.source_ids = [1]
            paragraph.evidence_quotes = {"1": "compares image representations"}
    with pytest.raises(ValidationFailed, match="процитировано 1"):
        validate_report(report)


def test_valid_report(report):
    validate_report(report)


def test_source_renumbering_follows_first_appearance(report):
    report.sections[0].paragraphs[0].source_ids = [15, 1]
    report.sections[0].paragraphs[0].evidence_quotes["15"] = "compares image representations"
    result = number_sources(report)
    assert result.sources[0].url.endswith("/15")
    assert result.sections[0].paragraphs[0].source_ids == [1, 2]
    assert result.sources[1].id == 2
    assert report.sources[0].id == 1


def test_csv_actual_computation_and_missing(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("x;y\n1;3\n2;4\n3;NaN\n;5\n", encoding="utf-8")
    stats = analyze_csv(path)
    assert stats["rows"] == 4
    assert stats["columns"]["x"]["mean"] == 2
    assert stats["columns"]["x"]["missing"] == 1
    assert stats["columns"]["y"]["non_numeric"] == 1
    assert len(stats["sha256"]) == 64


def test_owner_claim_atomic_and_persistent(settings):
    store = Store(settings.data_dir)
    assert not store.authorized(11, set())
    assert store.authorized(11, set(), claim=True)
    assert not store.authorized(12, set(), claim=True)
    assert Store(settings.data_dir).authorized(11, set())
    assert store.authorized(12, {12})


def test_restart_marks_job_retryable(settings):
    store = Store(settings.data_dir)
    store.save(11, {"stage": "running", "run_id": "persisted"})
    store.recover()
    assert store.get(11)["stage"] == "failed"
    assert store.get(11)["run_id"] == "persisted"


def test_faiss_search_includes_required_sources(settings, sources):
    corpus = Corpus(sources, settings)
    assert "15" in corpus.context("image representations", [15])


def test_bibliography_no_invented_pages(sources):
    text = bibliography(sources[0])
    assert "Электронный ресурс" in text and "URL:" in text
    assert "С. 1-" not in text


def test_settings_repr_hides_secrets():
    settings = Settings(token="secret-token", api_key="secret-key", scholar_key="other-key")
    assert "secret" not in repr(settings)


def test_scholar_api_metadata(settings):
    r = Retriever(settings)
    r.client.close()
    payload = {
        "data": [
            {
                "title": "Real metadata",
                "authors": [{"name": "Author"}],
                "year": 2025,
                "url": "https://semanticscholar.org/paper/test",
                "externalIds": {"DOI": "10.test/abc"},
                "abstract": "A real abstract " * 20,
                "venue": "Journal",
                "citationCount": 3,
            }
        ]
    }
    r.client = httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(200, json=payload)))
    source = r.semantic_scholar("test")[0]
    assert source["metadata_verified"] and source["doi"] == "10.test/abc"
    r.close()


def test_source_shortage_fails_closed(settings):
    r = Retriever(settings)
    for tool in r.tools.values():
        tool.func = lambda query: []
    r.arxiv = lambda query: []
    with pytest.raises(SearchUnavailable, match="Найдено 0"):
        r.collect(["test"])
    r.close()


def test_docx_layout_check_detects_wrong_margin(settings, report):
    from docx import Document
    from docx.shared import Mm

    from research_agent.reporter import Reporter

    reporter = Reporter(settings)
    chart = reporter.chart(report, settings.output_dir)
    path = settings.output_dir / "test.docx"
    reporter.document(report, path, chart)
    validate_docx(path)
    doc = Document(path)
    doc.sections[0].left_margin = Mm(20)
    doc.save(path)
    with pytest.raises(ValidationFailed, match="left_margin"):
        validate_docx(path)
